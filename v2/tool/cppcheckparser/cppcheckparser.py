import getopt
import os
import sys
import openpyxl
from openpyxl.utils import get_column_letter


IN_REPORT_FILE = 'cppcheck.txt'
OUT_REPORT_FILE = 'cppcheck.xlsx'

CODE_LOCATION_ROW = 1
CODE_LINE_ROW = 2
CODE_COMMITTER_ROW = 3
CODE_SEVERITY_ROW = 4
CODE_SUMMARY_ROW = 5

# cppcheck arguments:
#   --quiet --template='{file},{line},{severity},{id}:{message}'

# Input data format (separate with ',')
# code_path,line,severity,id+message

def open_and_parse_data(input_txt, output_xlsx):
    if not os.path.isfile(input_txt):
        raise BaseException("File not exist => %s" % (input_txt))

    with open(input_txt, "r") as infilestream:
        lines_data = infilestream.readlines()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "cppcheck"

    cell_row_index = 1
    ws.cell(cell_row_index, CODE_LOCATION_ROW).value = 'File Path'
    ws.column_dimensions[get_column_letter(CODE_LOCATION_ROW)].width = 32

    ws.cell(cell_row_index, CODE_LINE_ROW).value = 'Line'
    ws.column_dimensions[get_column_letter(CODE_LINE_ROW)].width = 7

    ws.cell(cell_row_index, CODE_COMMITTER_ROW).value = 'Committer'
    ws.column_dimensions[get_column_letter(CODE_COMMITTER_ROW)].width = 15

    ws.cell(cell_row_index, CODE_SEVERITY_ROW).value = 'Severity'
    ws.column_dimensions[get_column_letter(CODE_SEVERITY_ROW)].width = 15

    ws.cell(cell_row_index, CODE_SUMMARY_ROW).value = 'Summary'
    ws.column_dimensions[get_column_letter(CODE_SUMMARY_ROW)].width = 140

    cell_row_index = 2
    for line_data in lines_data:
        data_list = line_data.split(',', 4)
        system_command = 'git blame -p -L' + data_list[1] + ',' + \
                         data_list[1] + ' ' + data_list[0]
        system_out = os.popen(system_command)

        committer = 'Unknown'
        for out_line in system_out:
            if out_line.startswith('committer '):
                committer = out_line.split(' ', 1)[1]

        ws.cell(cell_row_index, CODE_LOCATION_ROW).value = data_list[0]
        ws.cell(cell_row_index, CODE_LINE_ROW).value = data_list[1]
        ws.cell(cell_row_index, CODE_COMMITTER_ROW).value = committer
        ws.cell(cell_row_index, CODE_SEVERITY_ROW).value = data_list[2]
        ws.cell(cell_row_index, CODE_SUMMARY_ROW).value = data_list[3]

        cell_row_index = cell_row_index + 1

    try:
        wb.save(output_xlsx)
    except Exception as err_msg:
        print("WARNING : Save excel file error")
        print("\t=> %s" % (str(err_msg)))


def parse_input_command(argv):
    param_dict = {
        'inFile': IN_REPORT_FILE,
        'outFile': OUT_REPORT_FILE,
    }
    if len(argv) == 0:
        return param_dict

    try:
        opts, _ = getopt.getopt(argv, "i:o:", ["input=", "output="])
    except getopt.GetoptError:
        print("cppcheckparser.py -i <input> -o <output>")
        sys.exit()

    for opt, arg in opts:
        if opt in ("-i", "--input"):
            param_dict['inFile'] = arg
        elif opt in ("-o", "--output"):
            param_dict['outFile'] = arg

    return param_dict


if __name__ == '__main__':
    param_dict = parse_input_command(sys.argv[1:])
    open_and_parse_data(param_dict['inFile'], param_dict['outFile'])
    print("Done")
